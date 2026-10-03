from __future__ import annotations

import copy
import io
import json
import random

import pytest
from docx import Document

from app.services.context_service import ContextService
from app.services.document_service import DocumentService


def _assert_coverage(result, sections, budget):
    originals = {
        s.get("index", i): s.get("content", s.get("text", "")) for i, s in enumerate(sections)
    }
    assert len(result["text"]) <= budget
    assert result["coverage"]["used_chars"] == len(result["text"])
    included = 0
    for ref in result["source_refs"]:
        source = originals[ref["section_index"]]
        assert ref["quote"] == source[ref["char_start"] : ref["char_end"]]
        assert ref["source_ref"] == (
            f"section:{ref['section_index']}:chars:{ref['char_start']}-{ref['char_end']}"
        )
        included += len(ref["quote"])
    coverage = result["coverage"]
    assert coverage["included_chars"] == included
    assert coverage["original_chars"] == coverage["included_chars"] + coverage["omitted_chars"]
    for section in coverage["sections"]:
        ranges = section["ranges"]
        assert ranges == sorted(ranges, key=lambda span: span["start"])
        assert all(a["end"] < b["start"] for a, b in zip(ranges, ranges[1:]))
        assert section["included_chars"] == sum(r["end"] - r["start"] for r in ranges)
    assert coverage["truncated"] == result["truncated"]
    json.dumps(result)


def test_complete_context_keeps_source_indices_and_verbatim_content():
    sections = [
        {
            "index": 4,
            "title": "Method",
            "content": "Encoder -> scoring -> output.",
            "page_start": 3,
        },
        {"index": 9, "title": "Results", "content": "Accuracy is 0.91; baseline is 0.75."},
    ]
    before = copy.deepcopy(sections)
    result = ContextService().select(sections, max_chars=1000)
    assert not result["truncated"]
    assert result["coverage"]["coverage_ratio"] == 1.0
    assert result["coverage"]["selected_section_indices"] == [4, 9]
    assert "Section 4: Method" in result["text"]
    assert "Section 9: Results" in result["text"]
    assert [s["content"] for s in result["sections"]] == [s["content"] for s in sections]
    assert result["sections"][0]["page_start"] == 3
    assert sections == before
    _assert_coverage(result, sections, 1000)


def test_selection_uses_document_order_not_request_order_or_duplicate_titles():
    sections = [{"title": "Repeated", "content": f"Source {i}"} for i in range(5)]
    result = ContextService().select(sections, [4, 1, 4], max_chars=1000)
    assert [s["index"] for s in result["sections"]] == [1, 4]
    assert result["coverage"]["unselected_section_indices"] == [0, 2, 3]
    assert not result["truncated"]
    _assert_coverage(result, sections, 1000)


def test_empty_selection_is_not_all_sections():
    result = ContextService().select([{"content": "Do not select this."}], [])
    assert result["text"] == ""
    assert result["sections"] == []
    assert result["coverage"]["selected_section_indices"] == []
    assert result["coverage"]["unselected_section_indices"] == [0]
    assert not result["truncated"]


def test_late_methods_results_and_table_survive_long_sections():
    sections = [
        {
            "index": 7,
            "title": "Paper",
            "content": (
                ("Background discussion without the key evidence. " * 700)
                + "\n\nOur proposed method uses an encoder, retrieval, and a gating objective.\n\n"
                + ("Additional descriptive material. " * 500)
                + "\n\nResults: accuracy improves to 0.91; the baseline achieves 0.75.\n\n"
                + "| Model | Accuracy |\n| --- | --- |\n| Baseline | 0.75 |\n| Ours | 0.91 |\n"
            ),
        }
    ]
    result = ContextService().select(sections, max_chars=2600)
    assert "encoder, retrieval, and a gating objective" in result["text"]
    assert "accuracy improves to 0.91" in result["text"]
    assert "| Baseline | 0.75 |\n| Ours | 0.91 |" in result["text"]
    assert "TRUNCATED CONTEXT" in result["text"]
    assert "source text omitted" in result["text"]
    assert result["truncated"]
    assert result["coverage"]["sections"][0]["status"] == "partial"
    _assert_coverage(result, sections, 2600)


def test_long_unbroken_section_samples_beyond_its_prefix():
    content = " ".join(f"token{i:05}" for i in range(10000))
    sections = [{"title": "Long section", "content": content}]
    result = ContextService().select(sections, max_chars=1600)
    refs = result["source_refs"]
    assert len(refs) >= 2
    assert min(r["char_start"] for r in refs) < 2000
    assert max(r["char_end"] for r in refs) > len(content) * 0.9
    assert result == ContextService().select(sections, max_chars=1600)
    _assert_coverage(result, sections, 1600)


def test_single_oversize_paragraph_does_not_only_keep_its_prefix():
    content = " ".join(f"term{i:03}" for i in range(100))
    sections = [{"title": "Paragraph", "content": content}]
    result = ContextService().select(sections, max_chars=500)
    assert any(ref["char_end"] > len(content) * 0.9 for ref in result["source_refs"])
    assert len(result["source_refs"]) >= 2
    _assert_coverage(result, sections, 500)


def test_budget_is_shared_and_short_sections_release_capacity():
    sections = [
        {"title": "Abstract", "content": "A short overview."},
        {"title": "Background", "content": "General observations. " * 1500},
        {
            "title": "Methods",
            "content": "The proposed architecture uses a gating objective. " * 1500,
        },
        {
            "title": "Results",
            "content": "Evaluation accuracy and baseline results are reported. " * 1500,
        },
    ]
    result = ContextService().select(sections, max_chars=6000)
    rows = result["coverage"]["sections"]
    assert [s["index"] for s in result["sections"]] == [0, 1, 2, 3]
    assert rows[0]["status"] == "full"
    assert rows[2]["included_chars"] > rows[1]["included_chars"]
    assert rows[3]["included_chars"] > rows[1]["included_chars"]
    assert result["coverage"]["used_chars"] > 5000
    _assert_coverage(result, sections, 6000)


def test_docx_table_context_is_exact_and_not_relocated_to_another_section():
    document = Document()
    document.add_heading("Methods", level=1)
    document.add_paragraph("An encoder produces the predictions. " * 700)
    document.add_heading("Results", level=1)
    document.add_paragraph("Experimental background. " * 1000)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Condition"
    table.cell(0, 1).text = "Score"
    table.cell(1, 0).text = "Measured condition"
    table.cell(1, 1).text = "82.7"
    stream = io.BytesIO()
    document.save(stream)
    sections = DocumentService().parse_docx(stream.getvalue())["sections"]

    result = ContextService().select(sections, [1], max_chars=1200)

    assert "Condition\tScore\nMeasured condition\t82.7" in result["text"]
    assert all(ref["section_index"] == 1 for ref in result["source_refs"])
    assert result["sections"][0]["source"]["file_type"] == "docx"
    _assert_coverage(result, sections, 1200)


def test_large_tables_keep_whole_rows_and_mark_omissions():
    rows = [f"Condition-{i:03}\t{i}.125\t{i}.875" for i in range(300)]
    sections = [{"title": "Results", "content": "\n".join(rows)}]
    result = ContextService().select(sections, max_chars=850)
    assert result["truncated"]
    assert result["source_refs"]
    for ref in result["source_refs"]:
        assert all(line in rows for line in ref["quote"].splitlines())
    assert any(ref["char_end"] > len(sections[0]["content"]) * 0.9 for ref in result["source_refs"])
    _assert_coverage(result, sections, 850)


def test_large_table_does_not_crowd_out_methods_and_results_in_same_section():
    rows = [f"Condition-{i:03}\t{i}.125\t{i}.875" for i in range(300)]
    content = (
        "Our method uses retrieval followed by a trained gate.\n\n"
        "Results: accuracy is 0.91 versus the baseline of 0.75.\n\n" + "\n".join(rows)
    )
    sections = [{"title": "Paper", "content": content}]
    result = ContextService().select(sections, max_chars=1500)
    assert "retrieval followed by a trained gate" in result["text"]
    assert "accuracy is 0.91 versus the baseline of 0.75" in result["text"]
    assert "Condition-" in result["text"]
    _assert_coverage(result, sections, 1500)


def test_key_evidence_inside_a_long_paragraph_is_not_lost_to_window_clipping():
    content = (
        "Earlier material. " * 900 + "Results: accuracy is 0.913 compared with baseline 0.721."
    )
    content += " Later material." * 900
    sections = [{"title": "Analysis", "content": content}]
    result = ContextService().select(sections, max_chars=550)
    assert "accuracy is 0.913 compared with baseline 0.721" in result["text"]
    _assert_coverage(result, sections, 550)


def test_thousands_of_short_table_rows_have_bounded_spread_selection():
    sections = [{"title": "Table", "content": "\n".join(f"{i}\t{i + 1}" for i in range(12000))}]
    result = ContextService().select(sections, max_chars=3200)
    assert result["source_refs"]
    assert max(ref["char_end"] for ref in result["source_refs"]) > len(sections[0]["content"]) * 0.9
    _assert_coverage(result, sections, 3200)


def test_unicode_source_offsets_are_character_offsets_not_bytes():
    sections = [{"title": "Unicode", "content": "\u03b1 \u03b2 \u7ed3\u679c = 0.91. " * 500}]
    result = ContextService().select(sections, max_chars=700)
    assert result["truncated"]
    _assert_coverage(result, sections, 700)


def test_deterministic_mixed_sections_keep_budget_and_exact_coverage():
    randomizer = random.Random(241)
    for _ in range(20):
        sections = [
            {
                "index": i * 2,
                "title": randomizer.choice(["Background", "Method", "Results"]),
                "content": "\n\n".join(
                    randomizer.choice(["Filler. ", "Results: score is 0.8. ", "Model\t0.75\n"])
                    * randomizer.randrange(1, 120)
                    for _ in range(randomizer.randrange(1, 5))
                ),
            }
            for i in range(randomizer.randrange(1, 6))
        ]
        budget = randomizer.randrange(0, 5000)
        before = copy.deepcopy(sections)
        result = ContextService().select(sections, max_chars=budget)
        _assert_coverage(result, sections, budget)
        assert result == ContextService().select(sections, max_chars=budget)
        assert sections == before


@pytest.mark.parametrize("budget", [0, 1, 10, 11, 50, 100, 200, 500, 1000, 32000])
def test_hard_budget_and_coverage_even_when_labels_do_not_fit(budget):
    sections = [
        {"title": "Very long title " * 50, "content": "Original material. " * 100},
        {"title": "Results", "content": "Observed accuracy is 0.83. " * 100},
    ]
    result = ContextService().select(sections, max_chars=budget)
    _assert_coverage(result, sections, budget)
    if budget <= 100:
        assert result["truncated"]
        assert result["coverage"]["omitted_section_indices"] == [0, 1]
    assert result == ContextService().select(sections, max_chars=budget)


@pytest.mark.parametrize("budget", [-1, 1.5, True, "1000"])
def test_invalid_budget_is_rejected(budget):
    with pytest.raises(ValueError, match="max_chars"):
        ContextService().select([], max_chars=budget)


@pytest.mark.parametrize("indices", [[9], [-1], [True], [0, False], ["0"]])
def test_invalid_selection_is_not_silently_replaced(indices):
    with pytest.raises(ValueError, match="indices"):
        ContextService().select([{"title": "Only section", "content": "Known."}], indices)


@pytest.mark.parametrize(
    "sections",
    [
        [{"index": 2}, {"index": 2}],
        [{"index": -1}],
        [{"index": True}],
        [{"index": 2, "section_index": 3}],
        [{"index": 0, "section_index": False}],
        [{"content": {"not": "text"}}],
    ],
)
def test_ambiguous_or_invalid_source_data_is_rejected(sections):
    with pytest.raises(ValueError):
        ContextService().select(sections)


def test_legacy_text_key_and_empty_sections_remain_supported():
    sections = [
        {"title": "Heading only", "content": ""},
        {"title": "Legacy", "text": "Legacy text."},
    ]
    result = ContextService().select(sections)
    assert not result["truncated"]
    assert [s["index"] for s in result["sections"]] == [0, 1]
    assert "Legacy text." in result["text"]
    _assert_coverage(result, sections, 32000)
