from __future__ import annotations

import io
import zipfile
from types import SimpleNamespace

import fitz
import pytest
from docx import Document

from app.core.exceptions import FileValidationException
from app.services import document_service
from app.services.document_service import DocumentService


def _docx_bytes(document) -> bytes:
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _rewrite_package(content: bytes, replacements: dict[str, bytes]) -> bytes:
    stream = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(content)) as original,
        zipfile.ZipFile(
            stream,
            "w",
            compression=zipfile.ZIP_DEFLATED,
        ) as updated,
    ):
        for name in original.namelist():
            updated.writestr(name, replacements.get(name, original.read(name)))
        for name, data in replacements.items():
            if name not in original.namelist():
                updated.writestr(name, data)
    return stream.getvalue()


@pytest.fixture
def service(monkeypatch):
    monkeypatch.setattr(
        document_service, "get_settings", lambda: SimpleNamespace(MAX_UPLOAD_SIZE_MB=1)
    )
    return DocumentService()


def test_docx_preserves_paragraphs_tables_and_source_order(service):
    document = Document()
    document.add_paragraph("Preamble before the first heading.")
    document.add_heading("Methods", level=1)
    document.add_paragraph("Input is encoded before scoring.")
    table = document.add_table(rows=3, cols=2)
    for row, values in zip(
        table.rows, [("Model", "Score"), ("Baseline", "0.75"), ("Ours", "0.91")]
    ):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.add_paragraph("Scores are averaged over three runs.")
    document.add_heading("Methods", level=2)
    document.add_paragraph("A distinct section with the same title.")
    content = _docx_bytes(document)

    parsed = service.parse(content, "docx")

    assert parsed == service.parse_docx(content)
    assert parsed["page_count"] is None
    assert parsed["full_text"] == (
        "Preamble before the first heading.\nMethods\nInput is encoded before scoring.\n"
        "Model\tScore\nBaseline\t0.75\nOurs\t0.91\nScores are averaged over three runs.\n"
        "Methods\nA distinct section with the same title."
    )
    assert [s["index"] for s in parsed["sections"]] == [0, 1, 2]
    assert [s["source_ref"] for s in parsed["sections"]] == ["section:0", "section:1", "section:2"]
    methods = parsed["sections"][1]
    assert methods["source"] == {
        "file_type": "docx",
        "section_index": 1,
        "heading_block_index": 1,
        "page_start": None,
        "page_end": None,
    }
    assert [b["type"] for b in methods["blocks"]] == ["paragraph", "table", "paragraph"]
    assert [b["block_index"] for b in methods["blocks"]] == [2, 3, 4]
    table_block = methods["blocks"][1]
    assert table_block["rows"] == [["Model", "Score"], ["Baseline", "0.75"], ["Ours", "0.91"]]
    assert methods["content"][table_block["start"] : table_block["end"]] == (
        "Model\tScore\nBaseline\t0.75\nOurs\t0.91"
    )
    for row, span in zip(table_block["rows"], table_block["row_spans"]):
        assert methods["content"][span["start"] : span["end"]] == "\t".join(row)


def test_docx_table_only_preserves_empty_cells_and_nested_order(service):
    document = Document()
    table = document.add_table(rows=1, cols=2)
    cell = table.cell(0, 0)
    cell.text = "Before nested table"
    nested = cell.add_table(rows=1, cols=2)
    nested.cell(0, 0).text = "Nested result"
    nested.cell(0, 1).text = "7.2"
    cell.add_paragraph("After nested table")

    parsed = service.parse_docx(_docx_bytes(document))

    section = parsed["sections"][0]
    assert section["title"] == "Untitled Section"
    assert section["content"] == "Before nested table\nNested result\t7.2\nAfter nested table\t"
    assert section["blocks"][0]["rows"] == [
        ["Before nested table\nNested result\t7.2\nAfter nested table", ""]
    ]
    assert section["blocks"][0]["end"] == len(section["content"])


def test_docx_merged_cells_do_not_duplicate_values(service):
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Shared heading"
    table.cell(1, 0).text = "Method A"
    table.cell(1, 1).text = "1.25"

    parsed = service.parse_docx(_docx_bytes(document))

    assert parsed["full_text"].count("Shared heading") == 1
    assert parsed["sections"][0]["blocks"][0]["rows"] == [["Shared heading"], ["Method A", "1.25"]]


@pytest.mark.parametrize("encoding", ["utf-8", "utf-16"])
def test_docx_rejects_dtd_without_resolving_entities(service, encoding):
    unsafe = (
        f'<?xml version="1.0" encoding="{encoding}"?>'
        '<!DOCTYPE document [<!ENTITY secret SYSTEM "file:///not-a-real-secret">]>'
        "<document>&secret;</document>"
    ).encode(encoding)
    content = _rewrite_package(_docx_bytes(Document()), {"word/document.xml": unsafe})

    with pytest.raises(FileValidationException, match="DTD"):
        service.parse_docx(content)
    with pytest.raises(FileValidationException, match="DTD"):
        service.validate_file("paper.docx", content, len(content))


def test_docx_checks_xml_parts_with_non_xml_extensions(service):
    content = _docx_bytes(Document())
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        types = archive.read("[Content_Types].xml").replace(
            b"</Types>",
            b'<Override PartName="/extra.data" ContentType="application/custom+xml"/></Types>',
        )
    content = _rewrite_package(
        content,
        {
            "[Content_Types].xml": types,
            "extra.data": b'<!DOCTYPE x [<!ENTITY a "expanded">]><x>&a;</x>',
        },
    )
    with pytest.raises(FileValidationException, match="DTD"):
        service.parse_docx(content)


def test_docx_accepts_xml_comments_without_treating_them_as_content_types(service):
    document = Document()
    document.add_paragraph("Retained text.")
    content = _docx_bytes(document)
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        types = archive.read("[Content_Types].xml").replace(
            b"</Types>", b"<!-- comment --></Types>"
        )
    parsed = service.parse_docx(_rewrite_package(content, {"[Content_Types].xml": types}))
    assert parsed["full_text"] == "Retained text."


@pytest.mark.parametrize("name", ["../escaped.xml", "/absolute.xml", "word\\bad.xml"])
def test_docx_rejects_unsafe_archive_paths(service, name):
    content = _rewrite_package(_docx_bytes(Document()), {name: b"<x/>"})
    with pytest.raises(FileValidationException, match="unsafe"):
        service.parse_docx(content)


def test_docx_rejects_duplicate_entries(service):
    stream = io.BytesIO(_docx_bytes(Document()))
    with zipfile.ZipFile(stream, "a") as archive, pytest.warns(UserWarning, match="Duplicate"):
        archive.writestr("word/document.xml", b"<document/>")
    with pytest.raises(FileValidationException, match="duplicate"):
        service.parse_docx(stream.getvalue())


@pytest.mark.parametrize(
    ("limit", "value", "message"),
    [
        ("_MAX_DOCX_MEMBERS", 2, "too many"),
        ("_MAX_DOCX_UNCOMPRESSED_BYTES", 1024, "expanded size"),
        ("_MAX_DOCX_XML_BYTES", 128, "XML part"),
    ],
)
def test_docx_expansion_is_bounded_before_loading(service, monkeypatch, limit, value, message):
    content = _docx_bytes(Document())
    monkeypatch.setattr(document_service, limit, value)
    monkeypatch.setattr(
        document_service, "DocxDocument", lambda *_: pytest.fail("Loaded unsafe ZIP")
    )
    with pytest.raises(FileValidationException, match=message):
        service.parse_docx(content)


def test_docx_rejects_invalid_xml_and_zip_disguised_as_docx(service):
    content = _rewrite_package(_docx_bytes(Document()), {"word/document.xml": b"<broken>"})
    with pytest.raises(FileValidationException, match="ZIP/XML"):
        service.parse(content, "docx")
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("unrelated.txt", "Not a DOCX document")
    with pytest.raises(FileValidationException, match="required document parts"):
        service.validate_file("paper.docx", stream.getvalue(), len(stream.getvalue()))
    with pytest.raises(FileValidationException, match="ZIP/XML"):
        service.parse_docx(b"PK\x03\x04not-a-zip")


def test_txt_indices_are_stable_for_preamble_and_repeated_headings(service):
    content = b"Preamble\n# Method\nFirst method.\n## Method\nSecond method."
    parsed = service.parse_txt(content)
    assert parsed["full_text"] == content.decode()
    assert [s["index"] for s in parsed["sections"]] == [0, 1, 2]
    assert [s["source"]["section_index"] for s in parsed["sections"]] == [0, 1, 2]
    assert parsed == service.parse(content, "txt")


def test_pdf_source_page_end_is_not_the_next_headings_page(service):
    with fitz.open() as document:
        for index in range(2):
            page = document.new_page()
            page.insert_text((72, 72), f"Heading {index}", fontsize=18)
            page.insert_text(
                (72, 100), "A sufficiently long body paragraph on this page.", fontsize=12
            )
        content = document.tobytes()
    parsed = service.parse_pdf(content)
    assert parsed["page_count"] == 2
    assert [(s["page_start"], s["page_end"]) for s in parsed["sections"]] == [(0, 0), (1, 1)]
    assert parsed["sections"][1]["source"]["page_start"] == 1


def test_empty_pdf_does_not_claim_ocr(service):
    with fitz.open() as document:
        document.new_page()
        content = document.tobytes()
    parsed = service.parse_pdf(content)
    assert parsed["sections"] == []
    assert parsed["full_text"] == ""
    assert parsed["page_count"] == 1
    assert parsed["metadata"]["ocr_performed"] is False
    assert "OCR is not performed" in parsed["warnings"][0]
    assert "ocr_markdown" not in parsed


def test_validate_file_preserves_interface_and_checks_actual_bytes(service):
    assert service.validate_file("PAPER.TXT", b"text", 4) == "txt"
    with pytest.raises(FileValidationException, match="Declared file size"):
        service.validate_file("paper.txt", b"text", 0)
    with pytest.raises(FileValidationException, match="exceeds"):
        service.validate_file("paper.txt", b"x" * (1024 * 1024 + 1), 1)
    with pytest.raises(FileValidationException, match="empty"):
        service.validate_file("paper.txt", b"", 0)
    with pytest.raises(FileValidationException, match="magic bytes"):
        service.validate_file("paper.pdf", b"text", 4)
    with pytest.raises(FileValidationException, match="Unsupported file extension"):
        service.validate_file("paper.exe", b"text", 4)
